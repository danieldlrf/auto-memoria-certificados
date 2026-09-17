import copy
from pathlib import Path
from typing import Dict
from docx import Document
from docx.oxml import OxmlElement
from docx.shared import Cm
from docx.table import Table
from models.student import Student


def _get_attr(student: Student, *attrs, default="") -> str:
    """Obtiene el primer atributo existente y no nulo del estudiante."""
    for attr in attrs:
        val = getattr(student, attr, None)
        if val is not None and str(val).strip() != "":
            return str(val)
    return default


def generate_memory(template_path: Path, students_by_dni: Dict[str, Student], output_path: Path) -> Path:
    """
    Genera la memoria del curso con python-docx directo.
    Rellena las tablas de calificaciones, tiempos de conexión, informes
    individuales y resumen de porcentajes alineados sin depender de Jinja2 ni docxtpl.
    """
    doc = Document(template_path)
    students = list(students_by_dni.values())

    # 1. PROCESAR TABLAS DE LISTADOS (Calificaciones y Tiempos de conexión)
    for table in doc.tables:
        header_text = "".join([c.text for c in table.rows[0].cells]).upper()

        # Tabla A: Calificación Final del Curso (7 columnas)
        if "EVALUAC" in header_text and len(table.columns) >= 5:
            while len(table.rows) > 1:
                tr = table.rows[-1]._tr
                tr.getparent().remove(tr)

            for student in students:
                row = table.add_row().cells
                full_name = _get_attr(student, 'full_name_v2', 'full_name', 
                                      default=f"{_get_attr(student, 'lastname')} {_get_attr(student, 'name')}".strip())
                row[0].text = full_name
                row[1].text = _get_attr(student, 'dni')
                row[2].text = _get_attr(student, 'ev1', default="-")
                row[3].text = _get_attr(student, 'ev2', default="-")
                row[4].text = _get_attr(student, 'evf', default="-")

        # Tabla B: Informe de Tiempo de Conexión (2 columnas)
        elif "TIEMPO TOTAL DE CONEXIÓN" in header_text or len(table.columns) == 2:
            if "CONEXIÓN" in header_text or "TIEMPO" in header_text:
                while len(table.rows) > 1:
                    tr = table.rows[-1]._tr
                    tr.getparent().remove(tr)

                for student in students:
                    row = table.add_row().cells
                    full_name = _get_attr(student, 'full_name_v2', 'full_name', 
                                          default=f"{_get_attr(student, 'lastname')} {_get_attr(student, 'name')}".strip())
                    row[0].text = full_name
                    row[1].text = _get_attr(student, 'time', 'time_connection', default="-")

    # 2. PROCESAR INFORME INDIVIDUALIZADO DE FORMACIÓN
    individual_table = None
    for table in doc.tables:
        table_text = "".join([c.text for r in table.rows for c in r.cells]).upper()
        if "EVALUACIÓN MÓDULO 1" in table_text or "EVALUACION MODULO 1" in table_text or "PRIMERA CONEXIÓN" in table_text:
            individual_table = table
            break

    if individual_table and students:
        clean_template_elem = copy.deepcopy(individual_table._element)
        current_elem = individual_table._element

        for idx, student in enumerate(students):
            if idx == 0:
                target_table = individual_table
            else:
                new_elem = copy.deepcopy(clean_template_elem)
                spacer_paragraph = OxmlElement('w:p')

                current_elem.addnext(spacer_paragraph)
                spacer_paragraph.addnext(new_elem)

                current_elem = new_elem
                target_table = Table(new_elem, doc)

            _rellenar_ficha_individual(target_table, student)

    # 3. PROCESAR RESUMEN DE PORCENTAJES ALINEADOS VERTICALMENTE
    _rellenar_resumen_estadisticas(doc, students)

    # 4. LIMPIEZA DE ETIQUETAS JINJA RESIDUALES EN PÁRRAFOS LIBRES
    for p in doc.paragraphs:
        if "{%" in p.text or "%}" in p.text:
            p.text = p.text.replace("{% for student in students %}", "").replace("{% endfor %}", "")

    output_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(output_path)
    return output_path


def _rellenar_ficha_individual(table: Table, student: Student) -> None:
    """Reemplaza los campos de texto dentro de una ficha individual de estudiante."""
    reemplazos = {
        "{{ student.lastname }}": _get_attr(student, 'lastname'),
        "{{ student.name }}": _get_attr(student, 'name'),
        "{{ student.dni }}": _get_attr(student, 'dni'),
        "{{ student.mail }}": _get_attr(student, 'mail', 'email'),
        "{{ student.movil }}": _get_attr(student, 'movil', 'phone', 'telephone'),
        "{{ student.firstconection }}": _get_attr(student, 'firstconection', 'first_connection'),
        "{{ student.time }}": _get_attr(student, 'time', default="-"),
        "{{ student.evf }}": _get_attr(student, 'evf', default="-"),
        "{{ student.ev1 }}": _get_attr(student, 'ev1', default="-"),
        "{{ student.ev2 }}": _get_attr(student, 'ev2', default="-"),
        "{{ student.ev3 }}": _get_attr(student, 'ev3', default="-"),
        "{{ student.ev4 }}": _get_attr(student, 'ev4', default="-"),
    }

    for row in table.rows:
        for cell in row.cells:
            for key, val in reemplazos.items():
                if key in cell.text:
                    cell.text = cell.text.replace(key, val)


def _rellenar_resumen_estadisticas(doc: Document, students: list) -> None:
    """Calcula los porcentajes de matriculados, aptos, no aptos y abandonos, y los alinea verticalmente."""
    total = len(students)
    aptos = 0
    no_aptos = 0
    abandonos = 0

    for s in students:
        evf_raw = _get_attr(s, 'evf')
        if not evf_raw or evf_raw.strip() in ("", "-"):
            abandonos += 1
        else:
            try:
                nota = float(str(evf_raw).replace(",", "."))
                if nota >= 5.0:
                    aptos += 1
                else:
                    no_aptos += 1
            except ValueError:
                val_str = str(evf_raw).strip().upper()
                if "APTO" in val_str and "NO" not in val_str:
                    aptos += 1
                elif "NO APTO" in val_str:
                    no_aptos += 1
                else:
                    abandonos += 1

    def _fmt_pct(val: float) -> str:
        return f"{int(val)}%" if val.is_integer() else f"{val:.1f}%"

    p_total = 100.0 if total > 0 else 0.0
    p_aptos = (aptos / total * 100.0) if total > 0 else 0.0
    p_no_aptos = (no_aptos / total * 100.0) if total > 0 else 0.0
    p_abandonos = (abandonos / total * 100.0) if total > 0 else 0.0

    # Obtener todos los párrafos del documento (libres e interiores de tablas)
    todos_parrafos = list(doc.paragraphs)
    for t in doc.tables:
        for r in t.rows:
            for c in r.cells:
                todos_parrafos.extend(c.paragraphs)

    for p in todos_parrafos:
        text_upper = p.text.upper()

        # 1. Eliminar completamente las líneas de puntos (…… o ......)
        if set(p.text.strip()).issubset({"…", ".", " "}) and len(p.text.strip()) > 0:
            p.text = ""
            continue

        # 2. Reemplazo por etiquetas Jinja si existen en la plantilla
        if "{{" in p.text:
            p.text = (p.text
                      .replace("{{ pct_matriculado }}", _fmt_pct(p_total))
                      .replace("{{ pct_apto }}", _fmt_pct(p_aptos))
                      .replace("{{ pct_no_apto }}", _fmt_pct(p_no_aptos))
                      .replace("{{ pct_abandonos }}", _fmt_pct(p_abandonos))
                      .replace("{{ total_matriculado }}", "")
                      .replace("{{ total_apto }}", "")
                      .replace("{{ total_no_apto }}", "")
                      .replace("{{ total_abandonos }}", ""))

        # 3. Reemplazo en texto plano alineando verticalmente mediante tabulación
        else:
            target_pct = None
            if "ALUMNADO MATRICULADO:" in text_upper:
                target_pct = _fmt_pct(p_total)
            elif "ALUMNADO APTO:" in text_upper:
                target_pct = _fmt_pct(p_aptos)
            elif "ALUMNADO NO APTO:" in text_upper:
                target_pct = _fmt_pct(p_no_aptos)
            elif "ABANDONOS:" in text_upper:
                target_pct = _fmt_pct(p_abandonos)

            if target_pct is not None:
                if ":" in p.text:
                    prefix = p.text.split(":")[0] + ":"
                else:
                    prefix = p.text.replace("%", "").strip()
                
                # Insertar tabulación y fijar parada de tabulador a 11 cm para alineación exacta
                p.text = f"{prefix}\t{target_pct}"
                p.paragraph_format.tab_stops.add_tab_stop(Cm(11))